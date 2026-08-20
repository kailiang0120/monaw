from __future__ import annotations

import json
import os
import hashlib
import shutil
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from app.agent.access_grant_broker import create_grant_ticket
from app.agent.approval_broker import create_ticket
from app.agent.controller_policy import ActionType, canonical, resolve_permission
from app.agent.execution_resume import register_executor
from app.agent.response_attachments import ATTACHMENT_HANDLE_PREFIX, resolve_attachment_handle_path

_MAX_READ_BYTES = 1024 * 1024
_DEFAULT_READ_BYTES = 65536
_MAX_LIST_LIMIT = 1000
_MAX_SEARCH_RESULTS = 200
_MAX_WRITE_BYTES = 16 * 1024 * 1024
_MAX_RECURSIVE_ITEMS = 5000
_MAX_RECURSIVE_BYTES = 512 * 1024 * 1024


@dataclass(frozen=True, slots=True)
class FilesystemTarget:
    raw: str
    resolved: Path
    parent: Path
    exists: bool
    is_symlink: bool

def _json(data: dict[str, Any]) -> str:
    return json.dumps(data, ensure_ascii=False)


def _error(reason_code: str, error: str, **extra: Any) -> str:
    return _json({"status": "error", "reason_code": reason_code, "error": error, **extra})

def _safe_int(value: Any, default: int, *, minimum: int, maximum: int) -> int:
    try:
        parsed = int(value)
    except (TypeError, ValueError):
        parsed = default
    return max(minimum, min(parsed, maximum))


def _iso_from_timestamp(value: float) -> str:
    return datetime.fromtimestamp(value, timezone.utc).isoformat()


def _stat_payload(path: Path) -> dict[str, Any]:
    stat = path.stat()
    return {
        "path": str(path),
        "exists": True,
        "type": "directory" if path.is_dir() else "file" if path.is_file() else "other",
        "is_file": path.is_file(),
        "is_dir": path.is_dir(),
        "size": stat.st_size,
        "modified_at": _iso_from_timestamp(stat.st_mtime),
        "created_at": _iso_from_timestamp(stat.st_ctime),
    }


def _missing_stat_payload(path: Path) -> dict[str, Any]:
    return {
        "path": str(path),
        "exists": False,
        "type": "missing",
        "is_file": False,
        "is_dir": False,
        "size": 0,
        "modified_at": "",
        "created_at": "",
    }


def _blocked_result(decision) -> str:
    return _json(
        {
            "status": "blocked",
            "reason": decision.reason,
            "reason_code": decision.reason_code,
            "policy_source": decision.policy_source,
        }
    )


def _pending_access_grant_result(
    decision,
    *,
    path: str,
    action_context: str,
    requested_access: str = "read",
) -> str:
    ticket = create_grant_ticket(
        target_type="path",
        target_identifier=path,
        display_name=path,
        action_context=action_context or decision.reason,
        requested_access=requested_access,
    )
    return _json(
        {
            "status": "pending_access_grant",
            "ticket_id": ticket.id,
            "target_type": ticket.target_type,
            "target_identifier": ticket.target_identifier,
            "display_name": ticket.display_name,
            "action_context": ticket.action_context,
            "requested_access": ticket.requested_access,
            "reason": decision.reason,
            "reason_code": decision.reason_code,
            "policy_source": decision.policy_source,
        }
    )


def _pending_approval_result(
    decision,
    *,
    tool_name: str,
    action_type: str,
    path: str,
    action_description: str,
    payload_args: dict[str, Any],
) -> str:
    input_str = json.dumps(payload_args, ensure_ascii=False, sort_keys=True)
    ticket = create_ticket(
        action_type=action_type,
        tool_name=tool_name,
        target_path=path,
        risk_level="medium",
        reason=decision.reason,
        action_description=action_description,
        payload={"input_str": input_str, "args": payload_args},
    )
    return _json(
        {
            "status": "pending_approval",
            "ticket_id": ticket.id,
            "action": action_description,
            "reason": decision.reason,
            "reason_code": decision.reason_code,
            "policy_source": decision.policy_source,
        }
    )


def _gate_path(
    action: ActionType,
    *,
    path: str,
    tool_name: str,
    action_description: str,
    payload_args: dict[str, Any],
) -> str | None:
    target_path = _policy_target_path(path)
    decision = resolve_permission(action, target_path=target_path)
    if decision.blocked:
        return _blocked_result(decision)
    if decision.requires_access_grant:
        requested_access = (
            "delete" if action == ActionType.DELETE
            else "write" if action in {ActionType.MUTATE, ActionType.EXEC}
            else "launch" if action == ActionType.LAUNCH_APP
            else "read"
        )
        return _pending_access_grant_result(
            decision,
            path=target_path,
            action_context=action_description,
            requested_access=requested_access,
        )
    if decision.requires_confirmation:
        return _pending_approval_result(
            decision,
            tool_name=tool_name,
            action_type=action.value,
            path=target_path,
            action_description=action_description,
            payload_args=payload_args,
        )
    return None


def _resolve_path(path: str) -> Path:
    attachment_path = resolve_attachment_handle_path(path)
    if attachment_path is not None:
        return attachment_path
    return Path(canonical(path))


def _policy_target_path(path: str) -> str:
    attachment_path = resolve_attachment_handle_path(path)
    return str(attachment_path) if attachment_path is not None else path


def _same_resolved_path(left: Path, right: Path) -> bool:
    return os.path.normcase(os.path.abspath(str(left))) == os.path.normcase(os.path.abspath(str(right)))


def _filesystem_target(path: str) -> FilesystemTarget:
    raw_path = Path(path).expanduser()
    resolved = _resolve_path(path)
    parent_source = raw_path.parent if str(raw_path.parent) else Path(".")
    return FilesystemTarget(
        raw=path,
        resolved=resolved,
        parent=_resolve_path(str(parent_source)),
        exists=resolved.exists(),
        is_symlink=raw_path.is_symlink() or resolved.is_symlink(),
    )


def _has_path_traversal(path: str) -> bool:
    return any(part == ".." for part in Path(path).parts)


def _validate_path_shape(path: str, *, operation: str) -> str | None:
    if str(path or "").strip().startswith(ATTACHMENT_HANDLE_PREFIX):
        if resolve_attachment_handle_path(path) is None:
            return _error("attachment_not_found", f"{operation} attachment handle is not available in this context.")
        return None
    if _has_path_traversal(path):
        return _error("path_traversal_rejected", f"{operation} path contains parent traversal segments.")
    return None


def _validate_write_budget(content: str, encoding: str) -> str | None:
    size = len(str(content or "").encode(encoding or "utf-8", errors="replace"))
    if size > _MAX_WRITE_BYTES:
        return _error(
            "write_size_limit_exceeded",
            "Write content exceeds filesystem write budget.",
            bytes=size,
            max_bytes=_MAX_WRITE_BYTES,
        )
    return None


def _tree_budget(path: Path) -> tuple[int, int, str | None]:
    if path.is_symlink():
        return 0, 0, f"Symlink or junction recursion is not allowed: {path}"
    if not path.is_dir():
        try:
            return 1, path.stat().st_size, None
        except OSError as exc:
            return 0, 0, str(exc)
    count = 0
    total = 0
    for item in path.rglob("*"):
        if item.is_symlink():
            return count, total, f"Symlink or junction recursion is not allowed: {item}"
        count += 1
        if count > _MAX_RECURSIVE_ITEMS:
            return count, total, "Recursive filesystem item limit exceeded."
        if item.is_file():
            try:
                total += item.stat().st_size
            except OSError as exc:
                return count, total, str(exc)
            if total > _MAX_RECURSIVE_BYTES:
                return count, total, "Recursive filesystem byte limit exceeded."
    return count, total, None


def _validate_recursive_budget(path: Path, *, operation: str) -> str | None:
    count, total, error = _tree_budget(path)
    if error:
        reason = "symlink_target_rejected" if "Symlink" in error else "recursive_budget_exceeded"
        return _error(
            reason,
            error,
            operation=operation,
            items=count,
            bytes=total,
            max_items=_MAX_RECURSIVE_ITEMS,
            max_bytes=_MAX_RECURSIVE_BYTES,
        )
    return None


def _revalidate_mutation_target(path: str) -> str | None:
    target = _filesystem_target(path)
    decision = resolve_permission(ActionType.MUTATE, target_path=str(target.parent if not target.exists else target.resolved))
    if decision.blocked:
        return _blocked_result(decision)
    if decision.requires_access_grant:
        return _pending_access_grant_result(
            decision,
            path=str(target.resolved),
            action_context="Revalidate filesystem mutation target",
            requested_access="write",
        )
    if target.is_symlink:
        return _error("symlink_target_rejected", f"Refusing to mutate a symlink or junction: {target.resolved}")
    return None


def file_stat(path: str) -> str:
    if not path:
        return _json({"status": "error", "error": "path is required"})
    shape_error = _validate_path_shape(path, operation="filesystem")
    if shape_error:
        return shape_error
    pending = _gate_path(
        ActionType.READ,
        path=path,
        tool_name="fs_stat",
        action_description=f"Stat file path: {path}",
        payload_args={"path": path},
    )
    if pending:
        return pending
    resolved = _resolve_path(path)
    stat = _stat_payload(resolved) if resolved.exists() else _missing_stat_payload(resolved)
    return _json({"status": "ok", "operation": "stat", "path": str(resolved), "exists": stat["exists"], "stat": stat})


def file_list(path: str, pattern: str = "*", recursive: bool = False, limit: int = 200) -> str:
    if not path:
        return _json({"status": "error", "error": "path is required"})
    shape_error = _validate_path_shape(path, operation="filesystem")
    if shape_error:
        return shape_error
    pending = _gate_path(
        ActionType.READ,
        path=path,
        tool_name="fs_list",
        action_description=f"List files under: {path}",
        payload_args={"path": path, "pattern": pattern, "recursive": recursive, "limit": limit},
    )
    if pending:
        return pending

    base = _resolve_path(path)
    if not base.exists():
        return _json({"status": "error", "error": f"Path does not exist: {base}"})
    if not base.is_dir():
        return _json({"status": "error", "error": f"Not a directory: {base}"})

    max_items = _safe_int(limit, 200, minimum=1, maximum=_MAX_LIST_LIMIT)
    glob_pattern = pattern or "*"
    iterator = base.rglob(glob_pattern) if recursive else base.glob(glob_pattern)
    entries: list[dict[str, Any]] = []
    for item in iterator:
        if len(entries) >= max_items:
            break
        try:
            entries.append(_stat_payload(item))
        except OSError:
            continue
    return _json({"status": "ok", "operation": "list", "path": str(base), "count": len(entries), "truncated": len(entries) >= max_items, "entries": entries})


def file_read(path: str, start: int = 0, max_bytes: int = _DEFAULT_READ_BYTES, encoding: str = "utf-8") -> str:
    if not path:
        return _json({"status": "error", "error": "path is required"})
    shape_error = _validate_path_shape(path, operation="filesystem")
    if shape_error:
        return shape_error
    pending = _gate_path(
        ActionType.READ,
        path=path,
        tool_name="fs_read",
        action_description=f"Read file: {path}",
        payload_args={"path": path, "start": start, "max_bytes": max_bytes, "encoding": encoding},
    )
    if pending:
        return pending

    resolved = _resolve_path(path)
    if not resolved.exists():
        return _json({"status": "error", "error": f"File does not exist: {resolved}"})
    if not resolved.is_file():
        return _json({"status": "error", "error": f"Not a file: {resolved}"})

    offset = _safe_int(start, 0, minimum=0, maximum=max(os.path.getsize(resolved), 0))
    byte_limit = _safe_int(max_bytes, _DEFAULT_READ_BYTES, minimum=1, maximum=_MAX_READ_BYTES)
    with resolved.open("rb") as handle:
        handle.seek(offset)
        data = handle.read(byte_limit)
    content = data.decode(encoding or "utf-8", errors="replace")
    next_start = offset + len(data)
    size = resolved.stat().st_size
    return _json(
        {
            "status": "ok",
            "operation": "read",
            "path": str(resolved),
            "start": offset,
            "bytes_read": len(data),
            "size": size,
            "truncated": next_start < size,
            "next_start": next_start if next_start < size else None,
            "content": content,
        }
    )


def file_write(
    path: str,
    content: str,
    overwrite: bool = False,
    create_parents: bool = True,
    encoding: str = "utf-8",
    dry_run: bool = False,
    *,
    _bypass_gate: bool = False,
) -> str:
    if not path:
        return _json({"status": "error", "error": "path is required"})
    shape_error = _validate_path_shape(path, operation="write")
    if shape_error:
        return shape_error
    budget_error = _validate_write_budget(content, encoding)
    if budget_error:
        return budget_error
    payload_args = {
        "path": path,
        "content": content,
        "overwrite": overwrite,
        "create_parents": create_parents,
        "encoding": encoding,
        "dry_run": dry_run,
    }
    if not _bypass_gate:
        pending = _gate_path(
            ActionType.MUTATE,
            path=path,
            tool_name="fs_write",
            action_description=f"Write file: {path}",
            payload_args=payload_args,
        )
        if pending:
            return pending

    revalidate = _revalidate_mutation_target(path)
    if revalidate:
        return revalidate
    resolved = _resolve_path(path)
    before = _stat_payload(resolved) if resolved.exists() else _missing_stat_payload(resolved)
    if resolved.exists() and resolved.is_dir():
        return _json({"status": "error", "error": f"Path is a directory: {resolved}"})
    if resolved.exists() and not overwrite:
        return _json({"status": "error", "error": f"File already exists: {resolved}", "reason_code": "file_exists"})
    if not resolved.parent.exists():
        if not create_parents:
            return _json({"status": "error", "error": f"Parent directory does not exist: {resolved.parent}"})
        if not dry_run:
            resolved.parent.mkdir(parents=True, exist_ok=True)
    if dry_run:
        return _json({"status": "ok", "operation": "write", "dry_run": True, "path": str(resolved), "before": before, "would_write_bytes": len(content.encode(encoding or "utf-8"))})

    with resolved.open("w", encoding=encoding or "utf-8", newline="") as handle:
        handle.write(content)
    return _json({"status": "ok", "operation": "write", "path": str(resolved), "before": before, "after": _stat_payload(resolved)})


def file_append(
    path: str,
    content: str,
    create_parents: bool = True,
    encoding: str = "utf-8",
    dry_run: bool = False,
    *,
    _bypass_gate: bool = False,
) -> str:
    if not path:
        return _json({"status": "error", "error": "path is required"})
    shape_error = _validate_path_shape(path, operation="append")
    if shape_error:
        return shape_error
    budget_error = _validate_write_budget(content, encoding)
    if budget_error:
        return budget_error
    payload_args = {
        "path": path,
        "content": content,
        "create_parents": create_parents,
        "encoding": encoding,
        "dry_run": dry_run,
    }
    if not _bypass_gate:
        pending = _gate_path(
            ActionType.MUTATE,
            path=path,
            tool_name="fs_append",
            action_description=f"Append file: {path}",
            payload_args=payload_args,
        )
        if pending:
            return pending

    revalidate = _revalidate_mutation_target(path)
    if revalidate:
        return revalidate
    resolved = _resolve_path(path)
    before = _stat_payload(resolved) if resolved.exists() else _missing_stat_payload(resolved)
    if resolved.exists() and resolved.is_dir():
        return _json({"status": "error", "error": f"Path is a directory: {resolved}"})
    if not resolved.parent.exists():
        if not create_parents:
            return _json({"status": "error", "error": f"Parent directory does not exist: {resolved.parent}"})
        if not dry_run:
            resolved.parent.mkdir(parents=True, exist_ok=True)
    if dry_run:
        return _json({"status": "ok", "operation": "append", "dry_run": True, "path": str(resolved), "before": before, "would_append_bytes": len(content.encode(encoding or "utf-8"))})

    with resolved.open("a", encoding=encoding or "utf-8", newline="") as handle:
        handle.write(content)
    return _json({"status": "ok", "operation": "append", "path": str(resolved), "before": before, "after": _stat_payload(resolved)})


def file_search(
    path: str,
    query: str,
    glob: str = "**/*",
    max_results: int = 50,
    case_sensitive: bool = False,
) -> str:
    if not path:
        return _json({"status": "error", "error": "path is required"})
    if not query:
        return _json({"status": "error", "error": "query is required"})
    shape_error = _validate_path_shape(path, operation="filesystem")
    if shape_error:
        return shape_error
    pending = _gate_path(
        ActionType.READ,
        path=path,
        tool_name="fs_search",
        action_description=f"Search files under: {path}",
        payload_args={
            "path": path,
            "query": query,
            "glob": glob,
            "max_results": max_results,
            "case_sensitive": case_sensitive,
        },
    )
    if pending:
        return pending

    base = _resolve_path(path)
    if not base.exists():
        return _json({"status": "error", "error": f"Path does not exist: {base}"})

    needle = query if case_sensitive else query.lower()
    limit = _safe_int(max_results, 50, minimum=1, maximum=_MAX_SEARCH_RESULTS)
    files = [base] if base.is_file() else [item for item in base.glob(glob or "**/*") if item.is_file()]
    matches: list[dict[str, Any]] = []
    for file_path in files:
        if len(matches) >= limit:
            break
        try:
            data = file_path.read_bytes()[:_DEFAULT_READ_BYTES]
            text = data.decode("utf-8", errors="replace")
        except OSError:
            continue
        haystack = text if case_sensitive else text.lower()
        if needle not in haystack:
            continue
        for line_number, line in enumerate(text.splitlines(), start=1):
            compare_line = line if case_sensitive else line.lower()
            column = compare_line.find(needle)
            if column < 0:
                continue
            matches.append(
                {
                    "path": str(file_path),
                    "line": line_number,
                    "column": column + 1,
                    "preview": line.strip()[:300],
                }
            )
            if len(matches) >= limit:
                break
    return _json({"status": "ok", "operation": "search", "path": str(base), "query": query, "count": len(matches), "truncated": len(matches) >= limit, "matches": matches})


def file_exists(path: str) -> str:
    if not path:
        return _json({"status": "error", "error": "path is required"})
    shape_error = _validate_path_shape(path, operation="filesystem")
    if shape_error:
        return shape_error
    pending = _gate_path(
        ActionType.READ,
        path=path,
        tool_name="fs_stat",
        action_description=f"Check file path: {path}",
        payload_args={"path": path},
    )
    if pending:
        return pending
    resolved = _resolve_path(path)
    stat = _stat_payload(resolved) if resolved.exists() else _missing_stat_payload(resolved)
    return _json({"status": "ok", "operation": "stat", "path": str(resolved), "exists": stat["exists"], "stat": stat})


def file_glob(path: str, pattern: str = "**/*", limit: int = 200) -> str:
    return file_list(path, pattern=pattern or "**/*", recursive=True, limit=limit)


def file_tree(path: str, depth: int = 2, limit: int = 200) -> str:
    if not path:
        return _json({"status": "error", "error": "path is required"})
    shape_error = _validate_path_shape(path, operation="filesystem")
    if shape_error:
        return shape_error
    pending = _gate_path(
        ActionType.READ,
        path=path,
        tool_name="fs_tree",
        action_description=f"Read file tree under: {path}",
        payload_args={"path": path, "depth": depth, "limit": limit},
    )
    if pending:
        return pending
    base = _resolve_path(path)
    if not base.exists():
        return _json({"status": "error", "error": f"Path does not exist: {base}"})
    if not base.is_dir():
        return _json({"status": "error", "error": f"Not a directory: {base}"})
    max_depth = _safe_int(depth, 2, minimum=0, maximum=8)
    max_items = _safe_int(limit, 200, minimum=1, maximum=_MAX_LIST_LIMIT)
    entries: list[dict[str, Any]] = []
    base_parts = len(base.parts)
    for item in base.rglob("*"):
        item_depth = len(item.parts) - base_parts
        if item_depth > max_depth:
            continue
        try:
            payload = _stat_payload(item)
        except OSError:
            continue
        payload["depth"] = item_depth
        payload["relative_path"] = str(item.relative_to(base))
        entries.append(payload)
        if len(entries) >= max_items:
            break
    return _json({"status": "ok", "operation": "tree", "path": str(base), "count": len(entries), "truncated": len(entries) >= max_items, "entries": entries})


def file_hash(path: str, algorithm: str = "sha256") -> str:
    if not path:
        return _json({"status": "error", "error": "path is required"})
    shape_error = _validate_path_shape(path, operation="filesystem")
    if shape_error:
        return shape_error
    pending = _gate_path(
        ActionType.READ,
        path=path,
        tool_name="fs_hash",
        action_description=f"Hash file: {path}",
        payload_args={"path": path, "algorithm": algorithm},
    )
    if pending:
        return pending
    resolved = _resolve_path(path)
    if not resolved.exists():
        return _json({"status": "error", "error": f"File does not exist: {resolved}"})
    if not resolved.is_file():
        return _json({"status": "error", "error": f"Not a file: {resolved}"})
    normalized = (algorithm or "sha256").lower()
    if normalized not in hashlib.algorithms_available:
        return _json({"status": "error", "error": f"Unsupported hash algorithm: {algorithm}"})
    digest = hashlib.new(normalized)
    with resolved.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return _json({"status": "ok", "operation": "hash", "path": str(resolved), "algorithm": normalized, "hash": digest.hexdigest(), "stat": _stat_payload(resolved)})


def file_patch(
    path: str,
    old: str,
    new: str,
    count: int = 1,
    encoding: str = "utf-8",
    dry_run: bool = False,
    *,
    _bypass_gate: bool = False,
) -> str:
    if not path:
        return _json({"status": "error", "error": "path is required"})
    if old == "":
        return _json({"status": "error", "error": "old text is required"})
    shape_error = _validate_path_shape(path, operation="patch")
    if shape_error:
        return shape_error
    budget_error = _validate_write_budget(new, encoding)
    if budget_error:
        return budget_error
    shape_error = _validate_path_shape(path, operation="patch")
    if shape_error:
        return shape_error
    budget_error = _validate_write_budget(new, encoding)
    if budget_error:
        return budget_error
    payload_args = {"path": path, "old": old, "new": new, "count": count, "encoding": encoding, "dry_run": dry_run}
    if not _bypass_gate:
        pending = _gate_path(
            ActionType.MUTATE,
            path=path,
            tool_name="fs_patch",
            action_description=f"Patch file: {path}",
            payload_args=payload_args,
        )
        if pending:
            return pending
    resolved = _resolve_path(path)
    if not resolved.exists():
        return _json({"status": "error", "error": f"File does not exist: {resolved}"})
    if not resolved.is_file():
        return _json({"status": "error", "error": f"Not a file: {resolved}"})
    before = _stat_payload(resolved)
    content = resolved.read_text(encoding=encoding or "utf-8")
    max_count = _safe_int(count, 1, minimum=1, maximum=100000)
    occurrences = content.count(old)
    if occurrences == 0:
        return _json({"status": "error", "error": "old text was not found", "reason_code": "patch_text_not_found"})
    updated = content.replace(old, new, max_count)
    replaced = min(occurrences, max_count)
    if dry_run:
        return _json({"status": "ok", "operation": "patch", "dry_run": True, "path": str(resolved), "replacements": replaced, "before": before})
    with resolved.open("w", encoding=encoding or "utf-8", newline="") as handle:
        handle.write(updated)
    return _json({"status": "ok", "operation": "patch", "path": str(resolved), "replacements": replaced, "before": before, "after": _stat_payload(resolved)})


def fs_mkdir(
    path: str,
    parents: bool = True,
    exist_ok: bool = True,
    dry_run: bool = False,
    *,
    _bypass_gate: bool = False,
) -> str:
    if not path:
        return _json({"status": "error", "error": "path is required"})
    shape_error = _validate_path_shape(path, operation="mkdir")
    if shape_error:
        return shape_error
    shape_error = _validate_path_shape(path, operation="mkdir")
    if shape_error:
        return shape_error
    payload_args = {"path": path, "parents": parents, "exist_ok": exist_ok, "dry_run": dry_run}
    if not _bypass_gate:
        pending = _gate_path(
            ActionType.MUTATE,
            path=path,
            tool_name="fs_mkdir",
            action_description=f"Create directory: {path}",
            payload_args=payload_args,
        )
        if pending:
            return pending
    resolved = _resolve_path(path)
    before = _stat_payload(resolved) if resolved.exists() else _missing_stat_payload(resolved)
    if resolved.exists() and not resolved.is_dir():
        return _json({"status": "error", "operation": "mkdir", "error": f"Path exists and is not a directory: {resolved}"})
    if resolved.exists() and not exist_ok:
        return _json({"status": "error", "operation": "mkdir", "error": f"Directory already exists: {resolved}", "reason_code": "directory_exists"})
    if dry_run:
        return _json({"status": "ok", "operation": "mkdir", "dry_run": True, "path": str(resolved), "before": before})
    resolved.mkdir(parents=parents, exist_ok=exist_ok)
    return _json({"status": "ok", "operation": "mkdir", "path": str(resolved), "before": before, "after": _stat_payload(resolved)})


def _gate_transfer_paths(
    *,
    operation: str,
    source: str,
    destination: str,
    source_action: ActionType,
    destination_action: ActionType,
    payload_args: dict[str, Any],
) -> str | None:
    source_pending = _gate_path(
        source_action,
        path=source,
        tool_name=f"fs_{operation}",
        action_description=f"{operation.title()} source: {source} -> {destination}",
        payload_args=payload_args,
    )
    if source_pending:
        return source_pending
    destination_pending = _gate_path(
        destination_action,
        path=destination,
        tool_name=f"fs_{operation}",
        action_description=f"{operation.title()} destination: {source} -> {destination}",
        payload_args=payload_args,
    )
    if destination_pending:
        return destination_pending
    return None


def fs_copy(
    source: str,
    destination: str,
    overwrite: bool = False,
    create_parents: bool = True,
    dry_run: bool = False,
    *,
    _bypass_gate: bool = False,
) -> str:
    if not source or not destination:
        return _json({"status": "error", "error": "source and destination are required"})
    for label, candidate in (("source", source), ("destination", destination)):
        shape_error = _validate_path_shape(candidate, operation=f"copy {label}")
        if shape_error:
            return shape_error
    payload_args = {
        "source": source,
        "destination": destination,
        "overwrite": overwrite,
        "create_parents": create_parents,
        "dry_run": dry_run,
    }
    if not _bypass_gate:
        pending = _gate_transfer_paths(
            operation="copy",
            source=source,
            destination=destination,
            source_action=ActionType.READ,
            destination_action=ActionType.MUTATE,
            payload_args=payload_args,
        )
        if pending:
            return pending
    revalidate = _revalidate_mutation_target(destination)
    if revalidate:
        return revalidate
    src = _resolve_path(source)
    dst = _resolve_path(destination)
    if not src.exists():
        return _json({"status": "error", "operation": "copy", "error": f"Source not found: {src}"})
    if _same_resolved_path(src, dst):
        stat = _stat_payload(src)
        return _json({"status": "ok", "operation": "copy", "source": str(src), "destination": str(dst), "same_path": True, "before": stat, "after": stat})
    budget_error = _validate_recursive_budget(src, operation="copy")
    if budget_error:
        return budget_error
    if dst.exists() and not overwrite:
        return _json({"status": "error", "operation": "copy", "error": f"Destination already exists: {dst}", "reason_code": "destination_exists"})
    before = _stat_payload(dst) if dst.exists() else _missing_stat_payload(dst)
    if dry_run:
        return _json({"status": "ok", "operation": "copy", "dry_run": True, "source": str(src), "destination": str(dst), "before": before})
    if not dst.parent.exists():
        if not create_parents:
            return _json({"status": "error", "operation": "copy", "error": f"Parent directory does not exist: {dst.parent}"})
        dst.parent.mkdir(parents=True, exist_ok=True)
    if src.is_dir():
        shutil.copytree(src, dst, dirs_exist_ok=overwrite)
    else:
        shutil.copy2(src, dst)
    return _json({"status": "ok", "operation": "copy", "source": str(src), "destination": str(dst), "before": before, "after": _stat_payload(dst)})


def fs_move(
    source: str,
    destination: str,
    overwrite: bool = False,
    create_parents: bool = True,
    dry_run: bool = False,
    *,
    _bypass_gate: bool = False,
) -> str:
    if not source or not destination:
        return _json({"status": "error", "error": "source and destination are required"})
    for label, candidate in (("source", source), ("destination", destination)):
        shape_error = _validate_path_shape(candidate, operation=f"move {label}")
        if shape_error:
            return shape_error
    payload_args = {
        "source": source,
        "destination": destination,
        "overwrite": overwrite,
        "create_parents": create_parents,
        "dry_run": dry_run,
    }
    if not _bypass_gate:
        pending = _gate_transfer_paths(
            operation="move",
            source=source,
            destination=destination,
            source_action=ActionType.MUTATE,
            destination_action=ActionType.MUTATE,
            payload_args=payload_args,
        )
        if pending:
            return pending
    revalidate = _revalidate_mutation_target(source) or _revalidate_mutation_target(destination)
    if revalidate:
        return revalidate
    src = _resolve_path(source)
    dst = _resolve_path(destination)
    if not src.exists():
        return _json({"status": "error", "operation": "move", "error": f"Source not found: {src}"})
    if _same_resolved_path(src, dst):
        stat = _stat_payload(src)
        return _json({"status": "ok", "operation": "move", "source": str(src), "destination": str(dst), "same_path": True, "before_source": stat, "before_destination": stat, "after": stat})
    budget_error = _validate_recursive_budget(src, operation="move")
    if budget_error:
        return budget_error
    if dst.exists() and not overwrite:
        return _json({"status": "error", "operation": "move", "error": f"Destination already exists: {dst}", "reason_code": "destination_exists"})
    before_source = _stat_payload(src)
    before_destination = _stat_payload(dst) if dst.exists() else _missing_stat_payload(dst)
    if dry_run:
        return _json({"status": "ok", "operation": "move", "dry_run": True, "source": str(src), "destination": str(dst), "before_source": before_source, "before_destination": before_destination})
    if not dst.parent.exists():
        if not create_parents:
            return _json({"status": "error", "operation": "move", "error": f"Parent directory does not exist: {dst.parent}"})
        dst.parent.mkdir(parents=True, exist_ok=True)
    if dst.exists() and overwrite:
        if dst.is_dir():
            shutil.rmtree(dst)
        else:
            dst.unlink()
    shutil.move(str(src), str(dst))
    return _json({"status": "ok", "operation": "move", "source": str(src), "destination": str(dst), "before_source": before_source, "before_destination": before_destination, "after": _stat_payload(dst)})


def fs_rename(path: str, new_name: str, overwrite: bool = False, dry_run: bool = False, *, _bypass_gate: bool = False) -> str:
    if not path or not new_name:
        return _json({"status": "error", "error": "path and new_name are required"})
    shape_error = _validate_path_shape(path, operation="rename")
    if shape_error:
        return shape_error
    if Path(new_name).name != new_name or _has_path_traversal(new_name):
        return _error("invalid_target_name", "new_name must be a single path segment.")
    source_label = path
    destination_label = str(Path(path).with_name(new_name))
    payload_args = {"path": path, "new_name": new_name, "overwrite": overwrite, "dry_run": dry_run}
    if not _bypass_gate:
        pending = _gate_transfer_paths(
            operation="rename",
            source=source_label,
            destination=destination_label,
            source_action=ActionType.MUTATE,
            destination_action=ActionType.MUTATE,
            payload_args=payload_args,
        )
        if pending:
            return pending
    revalidate = _revalidate_mutation_target(path)
    if revalidate:
        return revalidate
    source = _resolve_path(path)
    destination = source.with_name(new_name)
    if not source.exists():
        return _json({"status": "error", "operation": "rename", "error": f"Path not found: {source}"})
    if _same_resolved_path(source, destination):
        stat = _stat_payload(source)
        return _json({"status": "ok", "operation": "rename", "source": str(source), "destination": str(destination), "new_path": str(destination), "same_path": True, "before_source": stat, "before_destination": stat, "after": stat})
    if destination.exists() and not overwrite:
        return _json({"status": "error", "operation": "rename", "error": f"Destination already exists: {destination}", "reason_code": "destination_exists"})
    before_source = _stat_payload(source)
    before_destination = _stat_payload(destination) if destination.exists() else _missing_stat_payload(destination)
    if dry_run:
        return _json({"status": "ok", "operation": "rename", "dry_run": True, "source": str(source), "destination": str(destination), "before_source": before_source, "before_destination": before_destination})
    if destination.exists() and overwrite:
        if destination.is_dir():
            shutil.rmtree(destination)
        else:
            destination.unlink()
    source.rename(destination)
    return _json({"status": "ok", "operation": "rename", "source": str(source), "destination": str(destination), "new_path": str(destination), "before_source": before_source, "before_destination": before_destination, "after": _stat_payload(destination)})


def fs_delete(path: str, recursive: bool = True, dry_run: bool = False, *, _bypass_gate: bool = False) -> str:
    if not path:
        return _json({"status": "error", "error": "path is required"})
    shape_error = _validate_path_shape(path, operation="delete")
    if shape_error:
        return shape_error
    payload_args = {"path": path, "recursive": recursive, "dry_run": dry_run}
    if not _bypass_gate:
        pending = _gate_path(
            ActionType.DELETE,
            path=path,
            tool_name="fs_delete",
            action_description=f"Delete path: {path}",
            payload_args=payload_args,
        )
        if pending:
            return pending
    revalidate = _revalidate_mutation_target(path)
    if revalidate:
        return revalidate
    resolved = _resolve_path(path)
    if not resolved.exists():
        return _json({"status": "error", "operation": "delete", "error": f"Path not found: {resolved}"})
    budget_error = _validate_recursive_budget(resolved, operation="delete")
    if budget_error:
        return budget_error
    before = _stat_payload(resolved)
    if dry_run:
        return _json({"status": "ok", "operation": "delete", "dry_run": True, "path": str(resolved), "before": before})
    if resolved.is_dir():
        if not recursive:
            resolved.rmdir()
        else:
            shutil.rmtree(resolved)
    else:
        resolved.unlink()
    return _json({"status": "ok", "operation": "delete", "path": str(resolved), "before": before, "after": _missing_stat_payload(resolved)})


def create_file_compat(path: str, content: str = "", file_type: str = "text", *, _bypass_gate: bool = False) -> str:
    normalized_type = (file_type or "text").strip().lower()
    if normalized_type == "text":
        if not _bypass_gate:
            pending = _gate_path(
                ActionType.MUTATE,
                path=path,
                tool_name="create_file",
                action_description=f"Create text file: {path}",
                payload_args={"path": path, "content": content, "file_type": normalized_type},
            )
            if pending:
                return pending
        return file_write(path, content, overwrite=True, create_parents=True, _bypass_gate=True)

    payload_args = {"path": path, "content": content, "file_type": normalized_type}
    if not _bypass_gate:
        pending = _gate_path(
            ActionType.MUTATE,
            path=path,
            tool_name="create_file",
            action_description=f"Create {normalized_type} file: {path}",
            payload_args=payload_args,
        )
        if pending:
            return pending

    resolved = _resolve_path(path)
    before = _stat_payload(resolved) if resolved.exists() else _missing_stat_payload(resolved)
    resolved.parent.mkdir(parents=True, exist_ok=True)
    if normalized_type == "xlsx":
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.title = "Sheet1"
        if content:
            for row_idx, row in enumerate(content.strip().split("\n"), start=1):
                for col_idx, cell in enumerate(row.split(","), start=1):
                    ws.cell(row=row_idx, column=col_idx, value=cell.strip())
        wb.save(resolved)
    elif normalized_type == "docx":
        from docx import Document

        doc = Document()
        if content:
            doc.add_paragraph(content)
        doc.save(resolved)
    else:
        return _json({"status": "error", "error": f"Unsupported file_type '{file_type}'"})
    return _json({"status": "ok", "path": str(resolved), "file_type": normalized_type, "before": before, "after": _stat_payload(resolved)})


def read_file_compat(path: str) -> str:
    return file_read(path, max_bytes=8192)


def _parse_json_object(input_str: str) -> dict[str, Any]:
    try:
        value = json.loads(input_str) if input_str else {}
    except json.JSONDecodeError:
        return {}
    return value if isinstance(value, dict) else {}


def _resume_file_write(input_str: str) -> str:
    args = _parse_json_object(input_str)
    return file_write(
        str(args.get("path", "")),
        str(args.get("content", "")),
        bool(args.get("overwrite", False)),
        bool(args.get("create_parents", True)),
        str(args.get("encoding", "utf-8")),
        bool(args.get("dry_run", False)),
        _bypass_gate=True,
    )


def _resume_file_append(input_str: str) -> str:
    args = _parse_json_object(input_str)
    return file_append(
        str(args.get("path", "")),
        str(args.get("content", "")),
        bool(args.get("create_parents", True)),
        str(args.get("encoding", "utf-8")),
        bool(args.get("dry_run", False)),
        _bypass_gate=True,
    )


def _resume_create_file(input_str: str) -> str:
    args = _parse_json_object(input_str)
    return create_file_compat(
        str(args.get("path", "")),
        str(args.get("content", "")),
        str(args.get("file_type", "text")),
        _bypass_gate=True,
    )


def _resume_file_patch(input_str: str) -> str:
    args = _parse_json_object(input_str)
    return file_patch(
        str(args.get("path", "")),
        str(args.get("old", "")),
        str(args.get("new", "")),
        int(args.get("count", 1) or 1),
        str(args.get("encoding", "utf-8")),
        bool(args.get("dry_run", False)),
        _bypass_gate=True,
    )


def _resume_fs_mkdir(input_str: str) -> str:
    args = _parse_json_object(input_str)
    return fs_mkdir(
        str(args.get("path", "")),
        bool(args.get("parents", True)),
        bool(args.get("exist_ok", True)),
        bool(args.get("dry_run", False)),
        _bypass_gate=True,
    )


def _resume_fs_copy(input_str: str) -> str:
    args = _parse_json_object(input_str)
    return fs_copy(
        str(args.get("source", "")),
        str(args.get("destination", "")),
        bool(args.get("overwrite", False)),
        bool(args.get("create_parents", True)),
        bool(args.get("dry_run", False)),
        _bypass_gate=True,
    )


def _resume_fs_move(input_str: str) -> str:
    args = _parse_json_object(input_str)
    return fs_move(
        str(args.get("source", "")),
        str(args.get("destination", "")),
        bool(args.get("overwrite", False)),
        bool(args.get("create_parents", True)),
        bool(args.get("dry_run", False)),
        _bypass_gate=True,
    )


def _resume_fs_rename(input_str: str) -> str:
    args = _parse_json_object(input_str)
    return fs_rename(
        str(args.get("path", "")),
        str(args.get("new_name", "")),
        bool(args.get("overwrite", False)),
        bool(args.get("dry_run", False)),
        _bypass_gate=True,
    )


def _resume_fs_delete(input_str: str) -> str:
    args = _parse_json_object(input_str)
    return fs_delete(
        str(args.get("path", "")),
        bool(args.get("recursive", True)),
        bool(args.get("dry_run", False)),
        _bypass_gate=True,
    )


register_executor("fs_write", _resume_file_write)
register_executor("fs_append", _resume_file_append)
register_executor("fs_patch", _resume_file_patch)
register_executor("fs_mkdir", _resume_fs_mkdir)
register_executor("fs_copy", _resume_fs_copy)
register_executor("fs_move", _resume_fs_move)
register_executor("fs_rename", _resume_fs_rename)
register_executor("fs_delete", _resume_fs_delete)

# Legacy pending approvals may still exist after upgrade; keep these replay-only.
register_executor("file_write", _resume_file_write)
register_executor("file_append", _resume_file_append)
register_executor("create_file", _resume_create_file)
register_executor("file_patch", _resume_file_patch)

